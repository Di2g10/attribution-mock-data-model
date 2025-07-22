"""Global pytest fixtures for attrgen unit tests."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterator, Tuple

import openpyxl
import pytest

# ------------------------------------------------------------------

_OBJECTS_MINIMAL = [
    ("Name", "generate?", "row_count"),
    ("Company", "yes", 10),
]
_CHANNELS_MINIMAL = [
    ("Name",),
    ("Email",),
]
_ITYPES_MINIMAL = [
    ("Channel", "Delivered", "Opened"),
    ("Email", "Delivered", "Opened"),
]
_PARAMS_MINIMAL = [
    ("key", "value"),
    ("seed", 123),
]


@pytest.fixture(scope="session")
def workbook_path(tmp_path_factory: Any) -> Iterator[Path]:
    """Create a minimal Excel workbook compatible with Config loader."""
    tmp_dir = tmp_path_factory.mktemp("wb")
    wb_file = tmp_dir / "sample.xlsx"

    wb = openpyxl.Workbook()

    def _add(sheet_name: str, rows: list[Tuple[Any, ...]]) -> None:
        ws = wb.create_sheet(title=sheet_name)
        for r in rows:
            ws.append(list(r))

    _add("objects", _OBJECTS_MINIMAL)
    _add("Channels", _CHANNELS_MINIMAL)
    _add("Interaction Types", _ITYPES_MINIMAL)
    _add("params", _PARAMS_MINIMAL)

    # Remove the default sheet created by openpyxl
    wb.remove(wb[wb.sheetnames[0]])
    wb.save(wb_file)

    yield wb_file
